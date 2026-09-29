import os
import sys
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def apply_table_styling(table):
    """Apply standard table styling: header formatting, clean borders, cell margins, cantSplit."""
    tbl_elem = table._element
    tblPr = tbl_elem.find(f'{{{NS_W}}}tblPr')
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        tbl_elem.insert(0, tblPr)

    # Table Borders
    old_b = tblPr.find(f'{{{NS_W}}}tblBorders')
    if old_b is not None:
        tblPr.remove(old_b)

    new_borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(new_borders)

    # Cell Margins (Top/Bottom 120 dxa = 6pt, Left/Right 160 dxa = 8pt)
    old_mar = tblPr.find(f'{{{NS_W}}}tblCellMar')
    if old_mar is not None:
        tblPr.remove(old_mar)

    tblMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/>'
        f'<w:right w:w="160" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tblPr.append(tblMar)

    # Header Row
    tr_elems = tbl_elem.findall(f'{{{NS_W}}}tr')
    if tr_elems:
        hdr_tr = tr_elems[0]
        trPr = hdr_tr.find(f'{{{NS_W}}}trPr')
        if trPr is None:
            trPr = parse_xml(f'<w:trPr {nsdecls("w")}/>')
            hdr_tr.insert(0, trPr)
        if trPr.find(f'{{{NS_W}}}tblHeader') is None:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        if trPr.find(f'{{{NS_W}}}cantSplit') is None:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for tc in hdr_tr.findall(f'.//{{{NS_W}}}tc'):
            tcPr = tc.find(f'{{{NS_W}}}tcPr')
            if tcPr is None:
                tcPr = parse_xml(f'<w:tcPr {nsdecls("w")}/>')
                tc.insert(0, tcPr)
            shd = tcPr.find(f'{{{NS_W}}}shd')
            if shd is not None:
                tcPr.remove(shd)
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1E3A8A"/>'))

            for p in tc.findall(f'.//{{{NS_W}}}p'):
                pPr = p.find(f'{{{NS_W}}}pPr')
                if pPr is None:
                    pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
                    p.insert(0, pPr)
                jc = pPr.find(f'{{{NS_W}}}jc')
                if jc is not None:
                    pPr.remove(jc)
                pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))

                for r in p.findall(f'.//{{{NS_W}}}r'):
                    rPr = r.find(f'{{{NS_W}}}rPr')
                    if rPr is None:
                        rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
                        r.insert(0, rPr)
                    if rPr.find(f'{{{NS_W}}}b') is None:
                        rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
                    c = rPr.find(f'{{{NS_W}}}color')
                    if c is not None:
                        rPr.remove(c)
                    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="FFFFFF"/>'))

    # Data Rows
    for r_idx, tr in enumerate(tr_elems[1:]):
        trPr = tr.find(f'{{{NS_W}}}trPr')
        if trPr is None:
            trPr = parse_xml(f'<w:trPr {nsdecls("w")}/>')
            tr.insert(0, trPr)
        if trPr.find(f'{{{NS_W}}}cantSplit') is None:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
        for tc in tr.findall(f'.//{{{NS_W}}}tc'):
            tcPr = tc.find(f'{{{NS_W}}}tcPr')
            if tcPr is None:
                tcPr = parse_xml(f'<w:tcPr {nsdecls("w")}/>')
                tc.insert(0, tcPr)
            shd = tcPr.find(f'{{{NS_W}}}shd')
            if shd is not None:
                tcPr.remove(shd)
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{bg}"/>'))

def add_table_data(doc, headers, rows_data, keys_mapping):
    """
    Creates a styled table. If rows_data is empty, adds a single standard 'NIL' row.
    """
    col_count = len(headers)
    table = doc.add_table(rows=1, cols=col_count)
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

    valid_rows = []
    if rows_data and isinstance(rows_data, list):
        for item in rows_data:
            if not isinstance(item, dict):
                continue
            has_content = any(str(v).strip() for v in item.values() if v is not None)
            if has_content:
                valid_rows.append(item)

    if valid_rows:
        for idx, item in enumerate(valid_rows):
            row_cells = table.add_row().cells
            row_cells[0].text = str(idx + 1)
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p0.runs:
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(15, 23, 42)

            for col_idx, k in enumerate(keys_mapping):
                target_cell = row_cells[col_idx + 1]
                val = str(item.get(k, '') or '').strip()
                target_cell.text = val if val else '-'
                p = target_cell.paragraphs[0]
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(15, 23, 42)
    else:
        # Add single NIL row
        nil_row = table.add_row().cells
        for col_idx in range(col_count):
            nil_row[col_idx].text = "1" if col_idx == 0 else "NIL"
            p = nil_row[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(148, 163, 184)

    apply_table_styling(table)
    # Add normal spacing paragraph after table
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(2)
    p_space.paragraph_format.space_after = Pt(6)
    return len(valid_rows)

def generate_department_report(data, output_path):
    """
    Builds a DOCX report matching the EXACT format and fields of the existing monthly report.
    No extra fields, headers, banners, or signature blocks.
    """
    doc = docx.Document()

    # Standard Page Margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    dept = data.get('department', 'Civil Engineering')
    period = data.get('period', '01/04/2026 to 25/04/2026')
    hod_name = data.get('hodName', 'K Siva Prasad')
    submission_date = data.get('submissionDate', '25/04/2026')

    # P1: Exact Metadata Header Block (Matches template exactly!)
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(12)
    
    r1 = p_meta.add_run("Department: ")
    r1.bold = True
    p_meta.add_run(f"{dept}\n")
    
    r2 = p_meta.add_run("Reporting Period: ")
    r2.bold = True
    p_meta.add_run(f"{period}\n")
    
    r3 = p_meta.add_run("HOD Name: ")
    r3.bold = True
    p_meta.add_run(f"{hod_name}\n")
    
    r4 = p_meta.add_run("Date of Submission: ")
    r4.bold = True
    p_meta.add_run(f"{submission_date}")

    sections_data = data.get('sections', {})
    total_activities = 0

    def add_major_heading(text, desc=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(15, 23, 42)
        if desc:
            pd = doc.add_paragraph()
            pd.paragraph_format.space_before = Pt(0)
            pd.paragraph_format.space_after = Pt(6)
            rd = pd.add_run(desc)
            rd.font.size = Pt(9.5)
            rd.font.color.rgb = RGBColor(51, 65, 85)

    def add_sub_heading(text, desc=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(15, 23, 42)
        if desc:
            pd = doc.add_paragraph()
            pd.paragraph_format.space_before = Pt(0)
            pd.paragraph_format.space_after = Pt(5)
            rd = pd.add_run(desc)
            rd.font.size = Pt(9)
            rd.font.color.rgb = RGBColor(71, 85, 105)

    # -------------------------------------------------------------
    # 1. Research, Innovation & Entrepreneurship
    # -------------------------------------------------------------
    add_major_heading(
        "1. Research, Innovation & Entrepreneurship",
        "Include all faculty and departmental research output, innovative projects, and entrepreneurship initiatives relevant to your department. Attach links to reports/proofs wherever possible."
    )

    # a) Journal Publications (Table 0)
    add_sub_heading(
        "a) Journal Publications",
        "List all research articles, review papers, or technical notes published by faculty/students in peer-reviewed journals during the reporting period."
    )
    headers_0 = ['S. No.', 'Title of Publication', 'Authors(as mentioned in order)', 'Journal Name', 'ISSN/ISBN', 'Vol./Issue/Year', 'Page Nos.', 'Indexed In (e.g., Scopus)', 'Link to Publication/Document']
    keys_0 = ['title', 'authors', 'journalName', 'issnIsbn', 'volIssueYear', 'pageNos', 'indexedIn', 'link']
    total_activities += add_table_data(doc, headers_0, sections_data.get('journals', []), keys_0)

    # b) Conference Presentations (Table 1)
    add_sub_heading(
        "b) Conference Presentations",
        "Include papers presented at local/national/international conferences, symposiums, or workshops. Note presentation dates and attach link to presentation or conference proceedings."
    )
    headers_1 = ['S. No.', 'Title of Paper', 'Authors', 'Conference Name', 'Date', 'Location/Mode', 'Indexed In', 'Link to Presentation/Report']
    keys_1 = ['title', 'authors', 'conferenceName', 'date', 'locationMode', 'indexedIn', 'link']
    total_activities += add_table_data(doc, headers_1, sections_data.get('conferences', []), keys_1)

    # c) Patents (Table 2)
    add_sub_heading(
        "c) Patents",
        "Record granted or published patents, patent applications, and status updates for departmental innovations and intellectual property filings."
    )
    headers_2 = ['S. No.', 'Patent Title', 'Inventors(as per publication order)', 'Applicants', 'Patent Number', 'Patent Status (Filed/Published/Granted/Commercialized)', 'Awarded Date', 'Link to Document/Proof']
    keys_2 = ['title', 'inventors', 'applicants', 'patentNumber', 'status', 'awardedDate', 'link']
    total_activities += add_table_data(doc, headers_2, sections_data.get('patents', []), keys_2)

    # d) Entrepreneurship/Start-up Initiatives (Table 3)
    add_sub_heading(
        "d) Entrepreneurship/Start-up Initiatives",
        "List start-ups/spin-offs, business idea competitions, incubation activities, or innovation challenges in which the department/faculty/students participated."
    )
    headers_3 = ['S. No.', 'Entrepreneurship Activity/Program Title', 'Date', 'Type of Activity (Workshop/Competition/Incubation/Training/Mentorship etc.)', 'Participants (Student/Faculty/Alumni)', 'Organized By (Department/ED Cell/Incubator)', 'Mode (Online/Offline/Hybrid)', 'Key Outcomes (Startups Launched, Funding Received, Patents Filed, etc.)', 'No. of Participants', 'Mentor/Coordinator', 'Status (Ongoing/Completed)', 'Proof/Report Link']
    keys_3 = ['title', 'date', 'type', 'participants', 'organizedBy', 'mode', 'keyOutcomes', 'participantsCount', 'mentorCoordinator', 'status', 'link']
    total_activities += add_table_data(doc, headers_3, sections_data.get('entrepreneurship', []), keys_3)

    # -------------------------------------------------------------
    # 2. NSS and Other Extension Activities (Table 4)
    # -------------------------------------------------------------
    add_major_heading(
        "2. NSS and Other Extension Activities",
        "Capture all social outreach and community service work undertaken by the department, including NSS camps, awareness drives, and extension events."
    )
    headers_4 = ['S. No.', 'Event/Activity', 'Date', 'Venue', 'Type (NSS/Community)', 'No. of Participants', 'Type of Participants(Students/NSS Volunteers/Villagers/General Public/Faculty)', 'Outcomes', 'Coordinator', 'Report/Photo Link']
    keys_4 = ['event', 'date', 'venue', 'type', 'participantsCount', 'typeOfParticipants', 'outcomes', 'coordinator', 'link']
    total_activities += add_table_data(doc, headers_4, sections_data.get('nss', []), keys_4)

    # -------------------------------------------------------------
    # 3. Faculty Development Programs (FDPs) (Table 5)
    # -------------------------------------------------------------
    add_major_heading(
        "3. Faculty Development Programs (FDPs)",
        "List every short-term training, development course, workshop, or seminar attended or organized by faculty for professional development. Attach certificates where possible."
    )
    add_sub_heading(
        "a) Faculty Development Programs Attended",
        "Details of short-term training, courses, and workshops attended by faculty."
    )
    headers_5a = ['S. No.', 'Program Title', 'Type (FDP/Workshop/Seminar/Conference)', 'Dates', 'Organizing Body', 'Mode (Online/Offline/Hybrid)', 'Name of the faculty attended', 'Proof/Certificate Link']
    keys_5a = ['title', 'type', 'dates', 'organizingBody', 'mode', 'facultyAttended', 'link']
    fdp_att_data = sections_data.get('fdpAttended') or sections_data.get('fdp') or []
    total_activities += add_table_data(doc, headers_5a, fdp_att_data, keys_5a)

    add_sub_heading(
        "b) Faculty Development Programs Organized",
        "Details of faculty development programs and training organized by the department."
    )
    headers_5b = ['S. No.', 'Program Title', 'Type (FDP/Workshop/Seminar/Conference)', 'Dates', 'Dept. Organized', 'Mode (Online/Offline/Hybrid)', 'Resource person name, designation, co-organization and address', 'Name of the faculty coordinator/s', 'Proof/Certificate Link']
    keys_5b = ['title', 'type', 'dates', 'deptOrganized', 'mode', 'resourcePersonDetails', 'facultyCoordinators', 'link']
    fdp_org_data = sections_data.get('fdpOrganized') or []
    total_activities += add_table_data(doc, headers_5b, fdp_org_data, keys_5b)

    # -------------------------------------------------------------
    # 4. Student Development Programs (SDPs) (Table 6)
    # -------------------------------------------------------------
    add_major_heading(
        "4. Student Development Programs (SDPs)",
        "Include workshops, seminars, guest lectures, industrial visits, internships, symposiums, and community projects for students. Indicate type and outcomes for each."
    )
    headers_6 = ['S. No.', 'Event Title', 'Date', 'Type (Guest Lecture/Workshop/Seminar/Industrial Visit/Internship/Symposium/Community Project)', 'Resource Person with designation /Organization', 'Mode (Online/Offline/Hybrid)', 'Key Outcomes', 'No. of Participants', 'Coordinator', 'Report/Photo Link']
    keys_6 = ['title', 'date', 'type', 'resourcePerson', 'mode', 'keyOutcomes', 'participantsCount', 'coordinator', 'link']
    total_activities += add_table_data(doc, headers_6, sections_data.get('sdp', []), keys_6)

    # -------------------------------------------------------------
    # 5. Achievements & Awards
    # -------------------------------------------------------------
    add_major_heading(
        "5. Achievements & Awards",
        "Highlight significant recognitions, prizes, awards granted to faculty and students by academic, professional, or industry bodies."
    )
    # a) Faculty Achievements (Table 7)
    add_sub_heading("a) Faculty Achievements", "Recognitions for teaching, research, professional work, or leadership.")
    headers_7 = ['S. No.', 'Name', 'Award/Recognition', 'Organization/Body', 'Date', 'Proof/Certificate Link']
    keys_7 = ['name', 'award', 'organization', 'date', 'link']
    total_activities += add_table_data(doc, headers_7, sections_data.get('facultyAchievements', []), keys_7)

    # b) Student Achievements (Table 8)
    add_sub_heading("b) Student Achievements", "Achievements in academic, co-curricular, extra-curricular, or professional events.")
    headers_8 = ['S. No.', 'Name-Roll No', 'Award/Recognition', 'Event/Competition', 'Organization', 'Duration and date', 'Proof/Certificate Link']
    keys_8 = ['nameRoll', 'award', 'event', 'organization', 'durationDate', 'link']
    total_activities += add_table_data(doc, headers_8, sections_data.get('studentAchievements', []), keys_8)

    # c) Certifications (Table 9)
    add_sub_heading("c) Certifications", "")
    headers_9 = ['S. No.', 'Program Title', 'Type', 'Duration', 'Resource Person/Platform', 'Students/faculty Enrolled', 'Students/faculty  Certified', 'Key Outcomes', 'Evidence Link']
    keys_9 = ['title', 'type', 'duration', 'platform', 'enrolled', 'certified', 'keyOutcomes', 'link']
    total_activities += add_table_data(doc, headers_9, sections_data.get('certifications', []), keys_9)

    # -------------------------------------------------------------
    # 6. Other Notable Activities
    # -------------------------------------------------------------
    add_major_heading(
        "6. Other Notable Activities",
        "Any substantial department initiatives not covered above, such as meetings, collaborations, MoUs, or infrastructure changes."
    )
    # a) Department Meetings (Table 10)
    add_sub_heading("a) Department Meetings", "Details of official meetings: key decisions, date, and supporting documents.")
    headers_10 = ['S. No.', 'Date', 'Main Decisions/Topics Discussed', 'Policy Changes (if any)', 'Minutes/Proof Link']
    keys_10 = ['date', 'decisions', 'policyChanges', 'link']
    total_activities += add_table_data(doc, headers_10, sections_data.get('deptMeetings', []), keys_10)

    # b) Collaborations & MoUs (Table 11)
    add_sub_heading("b) Collaborations & MoUs", "Formal agreements/ongoing collaborations with industry, academia, or organizations.")
    headers_11 = ['S. No.', 'Name of Industry/Academic Body', 'Nature & Purpose', 'Date of Signing/Active Period', 'Faculty Involved(SPOC)', 'Supporting Documents/Link']
    keys_11 = ['name', 'purpose', 'datePeriod', 'facultySpoc', 'link']
    total_activities += add_table_data(doc, headers_11, sections_data.get('mous', []), keys_11)

    # -------------------------------------------------------------
    # 7. Additional/Other Relevant Initiatives (Table 12)
    # -------------------------------------------------------------
    add_major_heading(
        "7. Additional/Other Relevant Initiatives",
        "Alumni engagement, quality initiatives, special projects, or areas not elsewhere covered."
    )
    headers_12 = ['S. No.', 'Initiative/Activity', 'Date', 'Description', 'Outcomes', 'Coordinator(s)', 'Report/Link/Proof']
    keys_12 = ['initiative', 'date', 'description', 'outcomes', 'coordinator', 'link']
    total_activities += add_table_data(doc, headers_12, sections_data.get('additionalInitiatives', []), keys_12)

    # -------------------------------------------------------------
    # 8. Technical Association Activities (Table 13)
    # -------------------------------------------------------------
    add_major_heading(
        "8. Technical Association Activities",
        "Organizes technical workshops, competitions, industrial visits, and seminars to enhance technical skills."
    )
    headers_13 = ['S. No.', 'Event / Activity', 'Date', 'Type (Workshop/Seminar/Contest)', 'Resource Person / Coordinator', 'Participants', 'Outcomes / Achievements', 'Evidence / Proof Link']
    keys_13 = ['event', 'date', 'type', 'resourcePersonCoordinator', 'participants', 'outcomes', 'link']
    total_activities += add_table_data(doc, headers_13, sections_data.get('techAssociation', []), keys_13)

    # -------------------------------------------------------------
    # 9. IIC Cell (Institution’s Innovation Council) (Table 14)
    # -------------------------------------------------------------
    add_major_heading(
        "9. IIC Cell (Institution’s Innovation Council)",
        "Focuses on fostering innovation and entrepreneurship among students and faculty through various events and mentoring."
    )
    headers_14 = ['S. No.', 'Activity/Initiative', 'Date', 'Description/Objective', 'Resource Person/Partner', 'Beneficiaries', 'Key Outcomes/Impact', 'Evidence / Proof Link']
    keys_14 = ['activity', 'date', 'description', 'partner', 'beneficiaries', 'outcomes', 'link']
    total_activities += add_table_data(doc, headers_14, sections_data.get('iicCell', []), keys_14)

    # -------------------------------------------------------------
    # 10. Syllabus coverage Report (Summer Vacation holidays) (Table 15)
    # -------------------------------------------------------------
    add_major_heading("10. Syllabus coverage Report (Summer Vacation holidays)")
    headers_15 = ['S.No', 'Subject', 'Year/sem', 'Faculty', 'Syllabus Status - Completed', 'Syllabus Status - Pending', 'Remarks']
    keys_15 = ['subject', 'yearSem', 'faculty', 'completed', 'pending', 'remarks']
    total_activities += add_table_data(doc, headers_15, sections_data.get('syllabus', []), keys_15)

    # -------------------------------------------------------------
    # 11. Attendance shortage (Summer Vacation holidays)
    # Exact paragraph in template without table
    # -------------------------------------------------------------
    p_att = doc.add_paragraph()
    p_att.paragraph_format.space_before = Pt(12)
    p_att.paragraph_format.space_after = Pt(12)
    r_att = p_att.add_run("11. Attendance shortage (Summer Vacation holidays)")
    r_att.bold = True
    r_att.font.size = Pt(11)
    r_att.font.color.rgb = RGBColor(15, 23, 42)

    # Save Document
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    
    return {
        'output_path': output_path,
        'department': dept,
        'period': period,
        'hod_name': hod_name,
        'total_activities': total_activities
    }

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_department_report.py <input_json_file_or_string> <output_docx_path>")
        sys.exit(1)

    input_source = sys.argv[1]
    output_docx = sys.argv[2]

    if os.path.exists(input_source):
        with open(input_source, 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        data = json.loads(input_source)

    res = generate_department_report(data, output_docx)
    print(json.dumps(res))
