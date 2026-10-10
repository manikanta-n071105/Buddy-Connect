import os
import sys
import json
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

NS_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

DEPARTMENT_CODE_MAP = {
    'Civil Engineering': 'CIVIL',
    'Computer Science & Engineering': 'CSE',
    'Electronics & Communication Engineering': 'ECE',
    'Electrical & Electronics Engineering': 'EEE',
    'Mechanical Engineering': 'MECH',
    'Humanities & Sciences': 'H&S',
    'Innovation & Entrepreneurship': 'IIC/EDC',
    'Student Engagement & Clubs': 'CLUBS',
    'NSS & Community Engagement': 'NSS',
    'Minutes of the Meeting': 'MOM',
    'Training & Placement Cell': 'T&P',
    'Research & Development (R&D)': 'R&D',
    'Disciplinary Committee': 'DISCIP'
}

def get_dept_code(dept_name):
    return DEPARTMENT_CODE_MAP.get(dept_name, dept_name[:6].upper())

def apply_table_styling(table, col_widths=None):
    tbl_elem = table._element
    tblPr = tbl_elem.find(f'{{{NS_W}}}tblPr')
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        tbl_elem.insert(0, tblPr)

    # Table Borders (Clean Slate borders)
    old_b = tblPr.find(f'{{{NS_W}}}tblBorders')
    if old_b is not None:
        tblPr.remove(old_b)

    new_borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(new_borders)

    # Cell Margins
    old_mar = tblPr.find(f'{{{NS_W}}}tblCellMar')
    if old_mar is not None:
        tblPr.remove(old_mar)

    tblMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:left w:w="150" w:type="dxa"/>'
        f'<w:right w:w="150" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tblPr.append(tblMar)

    # Header Row Styling
    tr_elems = tbl_elem.findall(f'{{{NS_W}}}tr')
    if tr_elems:
        hdr_tr = tr_elems[0]
        trPr = hdr_tr.find(f'{{{NS_W}}}trPr')
        if trPr is None:
            trPr = parse_xml(f'<w:trPr {nsdecls("w")}/>')
            hdr_tr.insert(0, trPr)
        if trPr.find(f'{{{NS_W}}}tblHeader') is None:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        if trPr.find(f'{{{NS_W}}}cantSplit') is None:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for tc in hdr_tr.findall(f'.//{{{NS_W}}}tc'):
            tcPr = tc.find(f'{{{NS_W}}}tcPr')
            if tcPr is None:
                tcPr = parse_xml(f'<w:tcPr {nsdecls("w")}/>')
                tc.insert(0, tcPr)
            shd = tcPr.find(f'{{{NS_W}}}shd')
            if shd is not None:
                tcPr.remove(shd)
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1E3A8A"/>'))

            for p in tc.findall(f'.//{{{NS_W}}}p'):
                pPr = p.find(f'{{{NS_W}}}pPr')
                if pPr is None:
                    pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
                    p.insert(0, pPr)
                jc = pPr.find(f'{{{NS_W}}}jc')
                if jc is not None:
                    pPr.remove(jc)
                pPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))

                for r in p.findall(f'.//{{{NS_W}}}r'):
                    rPr = r.find(f'{{{NS_W}}}rPr')
                    if rPr is None:
                        rPr = parse_xml(f'<w:rPr {nsdecls("w")}/>')
                        r.insert(0, rPr)
                    if rPr.find(f'{{{NS_W}}}b') is None:
                        rPr.append(parse_xml(f'<w:b {nsdecls("w")}/>'))
                    c = rPr.find(f'{{{NS_W}}}color')
                    if c is not None:
                        rPr.remove(c)
                    rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="FFFFFF"/>'))

    # Data Rows Styling
    for r_idx, tr in enumerate(tr_elems[1:]):
        trPr = tr.find(f'{{{NS_W}}}trPr')
        if trPr is None:
            trPr = parse_xml(f'<w:trPr {nsdecls("w")}/>')
            tr.insert(0, trPr)
        if trPr.find(f'{{{NS_W}}}cantSplit') is None:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        bg = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
        for tc in tr.findall(f'.//{{{NS_W}}}tc'):
            tcPr = tc.find(f'{{{NS_W}}}tcPr')
            if tcPr is None:
                tcPr = parse_xml(f'<w:tcPr {nsdecls("w")}/>')
                tc.insert(0, tcPr)
            shd = tcPr.find(f'{{{NS_W}}}shd')
            if shd is not None:
                tcPr.remove(shd)
            tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{bg}"/>'))

    # Apply column widths if provided
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = w

def is_meaningful_item(item):
    if not item or not isinstance(item, dict):
        return False
    primary_keys = ['title', 'event', 'subject', 'initiative', 'name', 'nameRoll', 'decisions', 'award', 'platform', 'purpose', 'activity']
    for k in primary_keys:
        val = str(item.get(k, '') or '').strip()
        if len(val) >= 2:
            return True
    return any(len(str(v).strip()) >= 2 for k, v in item.items() if k not in ['status', 'mode', 'deptCode', 'noOfDays', 'link'])

def normalize_syllabus_year_sem(raw):
    if not raw: return '-'
    clean = str(raw).upper().strip()
    if '4-1' in clean or 'IV/I' in clean or 'IV-I' in clean or 'IV EEE' in clean or 'IV/1' in clean or '4/1' in clean:
        return 'IV / I'
    if '3-1' in clean or 'III/I' in clean or 'III-I' in clean or 'III EEE' in clean or 'III/1' in clean or '3/1' in clean:
        return 'III / I'
    if '2-1' in clean or 'II/I' in clean or 'II-I' in clean or 'II EEE' in clean or 'II/1' in clean or '2/1' in clean:
        return 'II / I'
    if '1-1' in clean or 'I/I' in clean or 'I-I' in clean or 'I/1' in clean or '1/1' in clean:
        return 'I / I'
    return str(raw).strip()

def normalize_syllabus_completed(completed, pending=None):
    if not completed: return '-'
    c = str(completed).strip()
    cUpper = c.upper()
    if '100%' in c: return '5 Units (100%)'
    if '95%' in c or '98%' in c: return '4.8 Units (95%)'
    if '90%' in c: return '4.5 Units (90%)'
    if '84%' in c: return '4.2 Units (84%)'
    if '80%' in c: return '4 Units (80%)'
    if '75%' in c: return '3.8 Units (75%)'
    if '70%' in c: return '3.5 Units (70%)'
    if '60%' in c: return '3 Units (60%)'
    if '50%' in c: return '2.5 Units (50%)'
    if '1,2,3,4-50%' in cUpper or '4-50%' in cUpper: return '3.5 Units'
    if '1,2,3,4' in cUpper or 'I,II,III,IV' in cUpper: return '4 Units'
    if '4 UNIT' in cUpper or c == '4' or '4 UNITS' in cUpper: return '4 Units'
    if '3.5' in cUpper or '3.5 UNITS' in cUpper: return '3.5 Units'
    if '1,2,3' in cUpper or 'I,II,III' in cUpper or c == '3' or '3 UNITS' in cUpper or '3 UNIT' in cUpper: return '3 Units'
    if 'I,II' in cUpper or '1,2' in cUpper or c == '2' or '2 UNITS' in cUpper or '2 UNIT' in cUpper: return '2 Units'
    if '1 UNIT' in cUpper or c == '1': return '1 Unit'
    if 'EXPERIMENT' in cUpper: return c
    if 'GENERAL TOPICS' in cUpper: return '100% Completed'
    return c

def normalize_syllabus_pending(completed, pending):
    p = str(pending or '').strip()
    pUpper = p.upper()
    cUpper = str(completed or '').upper().strip()
    if '5-80%' in pUpper or '5-60%' in pUpper or '5-50%' in pUpper or '5 STARTED' in pUpper:
        return 'Unit 5 (In Progress)'
    if '5 UNIT' in pUpper or p == '5' or 'V UNIT' in pUpper or '5TH UNIT' in pUpper or '5 UNIT RUNNING' in pUpper:
        return 'Unit 5'
    if '4,5' in pUpper or '4, 5' in pUpper or 'IV,V' in pUpper or '4 UNIT AND 5' in pUpper or '4 UNIT RUNNING' in pUpper or '4 RUNNING' in pUpper or 'IV RUNNING' in pUpper or 'IV STARTED' in pUpper:
        return 'Units 4 & 5'
    if '3 UNIT RUNNING' in pUpper or '3 RUNNING' in pUpper or 'III RUNNING' in pUpper or 'III STARTED' in pUpper or 'III UNIT 30%' in pUpper:
        return 'Units 3, 4 & 5'
    if 'IV UNIT' in pUpper or p == '4':
        return 'Units 4 & 5'
    if not p or p == '-' or p == 'NIL':
        if cUpper == '4' or '4 UNITS' in cUpper or '1,2,3,4' in cUpper or '80%' in cUpper:
            return 'Unit 5'
        if cUpper == '3.5' or '3.5 UNITS' in cUpper or '70%' in cUpper:
            return 'Unit 4 (50%) & Unit 5'
        if cUpper == '3' or '3 UNITS' in cUpper or '1,2,3' in cUpper or '60%' in cUpper:
            return 'Units 4 & 5'
        if cUpper == '2' or '2 UNITS' in cUpper or '1,2' in cUpper or '50%' in cUpper:
            return 'Units 3, 4 & 5'
        if '100%' in cUpper or cUpper == '5' or 'GENERAL TOPICS' in cUpper:
            return 'Nil (Complete)'
        return '-'
    return p

def normalize_syllabus_remarks(remarks, completed=None):
    r = str(remarks or '').strip()
    if not r or r == '-' or r == 'NIL':
        return 'On Schedule as per Lesson Plan'
    return r

def add_styled_cell_text(cell, lines, font_size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Adds multiple formatted text lines/runs to a table cell."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15

    for i, line in enumerate(lines):
        if not line:
            continue
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15

        prefix = ""
        text = ""
        is_bold = False
        color_hex = None

        if isinstance(line, tuple) or isinstance(line, list):
            if len(line) == 4:
                prefix, text, is_bold, color_hex = line
            elif len(line) == 3:
                text, is_bold, color_hex = line
            elif len(line) == 2:
                text, is_bold = line
            elif len(line) == 1:
                text = line[0]
        elif isinstance(line, str):
            text = line

        if prefix:
            r_pre = p.add_run(str(prefix))
            r_pre.bold = True
            r_pre.font.size = Pt(font_size)
            r_pre.font.color.rgb = RGBColor(30, 58, 138)
        
        if text:
            r = p.add_run(str(text))
            r.bold = bool(is_bold)
            r.font.size = Pt(font_size)
            if color_hex:
                color_hex = str(color_hex).lstrip('#')
                try:
                    r.font.color.rgb = RGBColor(int(color_hex[:2], 16), int(color_hex[2:4], 16), int(color_hex[4:], 16))
                except Exception:
                    pass
            else:
                r.font.color.rgb = RGBColor(15, 23, 42)

def generate_consolidated_master_docx(all_dept_data, period, output_path, college_name="SANSKRITHI SCHOOL OF ENGINEERING"):
    doc = docx.Document()

    # Standard Page Margins: 0.65 in
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # 1. Master Header Banner
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_col = p_title.add_run(f"{college_name}\n")
    r_col.bold = True
    r_col.font.size = Pt(13.5)
    r_col.font.color.rgb = RGBColor(30, 58, 138)

    r_sub = p_title.add_run(f"CONSOLIDATED INSTITUTIONAL PROGRESS REPORT — {period.upper()}\n")
    r_sub.bold = True
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(194, 65, 12)

    unique_depts = list(all_dept_data.keys())
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(4)
    p_meta.paragraph_format.space_after = Pt(10)
    r_meta = p_meta.add_run(f"Reporting Period: {period}  |  Participating Entities: {', '.join(unique_depts)}")
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(71, 85, 105)

    def add_major_heading(text, desc=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(15, 23, 42)
        if desc:
            pd = doc.add_paragraph()
            pd.paragraph_format.space_before = Pt(0)
            pd.paragraph_format.space_after = Pt(4)
            rd = pd.add_run(desc)
            rd.font.size = Pt(8.5)
            rd.font.color.rgb = RGBColor(71, 85, 105)

    # Helper to gather all items for a section key across all departments
    def gather_items(sec_key, alt_key=None):
        res = []
        for d_name, sec_data in all_dept_data.items():
            if not isinstance(sec_data, dict):
                continue
            items = sec_data.get(sec_key, [])
            if not items and alt_key:
                items = sec_data.get(alt_key, [])
            if isinstance(items, list):
                for it in items:
                    if is_meaningful_item(it):
                        it_copy = dict(it)
                        it_copy['department'] = d_name
                        it_copy['deptCode'] = get_dept_code(d_name)
                        res.append(it_copy)
        return res

    # -------------------------------------------------------------
    # Executive Summary Matrix Table (Exact same as PDF)
    # -------------------------------------------------------------
    add_major_heading("INSTITUTIONAL EXECUTIVE SUMMARY MATRIX", "Cross-departmental consolidated performance metrics across all academic and research domains.")
    
    matrix_headers = ['Domain / Metric Category', 'CIVIL', 'CSE', 'ECE', 'EEE', 'MECH', 'H&S', 'TOTAL']
    matrix_table = doc.add_table(rows=1, cols=len(matrix_headers))
    matrix_table.autofit = False

    for i, h in enumerate(matrix_headers):
        matrix_table.rows[0].cells[i].text = h
        p = matrix_table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(255, 255, 255)

    academic_codes = ['CIVIL', 'CSE', 'ECE', 'EEE', 'MECH', 'H&S']
    
    metric_categories = [
        ('1a. Journal Publications', 'journals'),
        ('1b. Conference Presentations', 'conferences'),
        ('2a. Patents', 'patents'),
        ('2b. Activities and Initiatives', 'entrepreneurship'),
        ('3a. FDPs Attended', 'fdpAttended'),
        ('3b. FDPs Organized', 'fdpOrganized'),
        ('4. Student Development Programs (SDPs)', 'sdp'),
        ('5a. Faculty Achievements', 'facultyAchievements'),
        ('5b. Student Achievements', 'studentAchievements'),
        ('5c. Certifications', 'certifications'),
        ('6. Collaborations & MoUs', 'mous'),
        ('7. Technical Association Activities', 'techAssociation'),
        ('8. Syllabus Coverage Reports', 'syllabus'),
        ('9. Clubs & Student Engagement Activity', 'studentEngagement'),
        ('10. NSS and Other Extension Activities', 'nss'),
        ('11. Additional/Other Relevant Initiatives', 'additionalInitiatives'),
    ]

    col_totals = {c: 0 for c in academic_codes}
    grand_total = 0

    for label, sec_key in metric_categories:
        row_cells = matrix_table.add_row().cells
        row_cells[0].text = label
        p0 = row_cells[0].paragraphs[0]
        for r in p0.runs:
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(15, 23, 42)

        row_sum = 0
        for idx, code in enumerate(academic_codes):
            matching_sec = None
            for d_name, sec_data in all_dept_data.items():
                if get_dept_code(d_name) == code or d_name.upper().startswith(code):
                    matching_sec = sec_data
                    break
            
            cnt = 0
            if matching_sec and isinstance(matching_sec, dict):
                items = matching_sec.get(sec_key, [])
                if sec_key == 'fdpAttended' and not items:
                    items = matching_sec.get('fdp', [])
                cnt = len([it for it in items if is_meaningful_item(it)])

            col_totals[code] += cnt
            row_sum += cnt
            cell = row_cells[idx + 1]
            cell.text = str(cnt) if cnt > 0 else '-'
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(8.5)
                if cnt > 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(30, 58, 138)
                else:
                    r.font.color.rgb = RGBColor(148, 163, 184)

        grand_total += row_sum
        tot_cell = row_cells[-1]
        tot_cell.text = str(row_sum)
        pt = tot_cell.paragraphs[0]
        pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in pt.runs:
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(194, 65, 12)

    # Total row
    tot_row = matrix_table.add_row().cells
    tot_row[0].text = "OVERALL TOTAL METRICS"
    p_tot_lbl = tot_row[0].paragraphs[0]
    for r in p_tot_lbl.runs:
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(30, 58, 138)

    for idx, code in enumerate(academic_codes):
        c_cell = tot_row[idx + 1]
        c_cell.text = str(col_totals[code])
        p = c_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(30, 58, 138)

    g_cell = tot_row[-1]
    g_cell.text = str(grand_total)
    p_g = g_cell.paragraphs[0]
    p_g.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p_g.runs:
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(194, 65, 12)

    apply_table_styling(matrix_table, [Inches(2.5), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.65), Inches(0.75)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1A. Journal Publications (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("1A. JOURNAL PUBLICATIONS", "Peer-reviewed research articles published in indexed journals.")
    j_headers = ['S.No', 'Branch', 'Paper Title', 'Authors & Publication Details', 'Indexing & Verification']
    j_table = doc.add_table(rows=1, cols=len(j_headers))
    j_table.autofit = False
    for i, h in enumerate(j_headers):
        j_table.rows[0].cells[i].text = h
    
    j_items = gather_items('journals')
    if j_items:
        for idx, it in enumerate(j_items):
            row = j_table.add_row().cells
            row[0].text = str(idx + 1)
            row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']
            row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', '')], font_size=8.5)
            
            auth_lines = [
                ("Authors: ", it.get('authors', '-'), False, "334155"),
                ("Journal: ", f"{it.get('journalName', '')} ({it.get('volIssueYear', '')})", False, "475569"),
                ("Pages: ", it.get('pageNos', '-'), False, "64748B")
            ]
            add_styled_cell_text(row[3], auth_lines, font_size=8)

            idx_lines = [
                ("ISSN: ", it.get('issnIsbn', '-'), False, "334155"),
                ("Indexed: ", it.get('indexedIn', '-'), True, "1E3A8A"),
            ]
            if it.get('link'):
                idx_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[4], idx_lines, font_size=8)
    else:
        nil = j_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(j_table, [Inches(0.4), Inches(0.65), Inches(2.6), Inches(2.1), Inches(1.45)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1B. Conference Presentations (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("1B. CONFERENCE PRESENTATIONS", "Papers presented at national and international conferences.")
    c_headers = ['S.No', 'Branch', 'Presentation Title', 'Authors & Conference Details', 'Date, Venue & Indexing']
    c_table = doc.add_table(rows=1, cols=len(c_headers))
    for i, h in enumerate(c_headers): c_table.rows[0].cells[i].text = h

    c_items = gather_items('conferences')
    if c_items:
        for idx, it in enumerate(c_items):
            row = c_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', '')], font_size=8.5)
            
            c_lines = [
                ("Authors: ", it.get('authors', '-'), False, "334155"),
                ("Conference: ", it.get('conferenceName', '-'), False, "475569")
            ]
            if it.get('volIssueYear') or it.get('pageNos'):
                c_lines.append(("", f"Issue: {it.get('volIssueYear', '-')} | Pages: {it.get('pageNos', '-')}", False, "64748B"))
            add_styled_cell_text(row[3], c_lines, font_size=8)

            d_lines = [
                ("Date: ", it.get('date', '-'), False, "334155"),
                ("Venue: ", it.get('locationMode', '-'), False, "475569"),
                ("Indexed: ", it.get('indexedIn', '-'), True, "1E3A8A")
            ]
            if it.get('link'): d_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[4], d_lines, font_size=8)
    else:
        nil = c_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(c_table, [Inches(0.4), Inches(0.65), Inches(2.5), Inches(2.2), Inches(1.45)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 2A. Patents (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("2A. PATENTS", "Granted or published patents and official patent filings.")
    p_headers = ['S.No', 'Branch', 'Patent / Innovation Title', 'Inventors & Registration', 'Status & Date']
    p_table = doc.add_table(rows=1, cols=len(p_headers))
    for i, h in enumerate(p_headers): p_table.rows[0].cells[i].text = h

    p_items = gather_items('patents')
    if p_items:
        for idx, it in enumerate(p_items):
            row = p_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', '')], font_size=8.5)
            add_styled_cell_text(row[3], [("Inv: ", it.get('inventors', '-'), False, "334155"), ("Pat: ", it.get('patentNumber', '-'), False, "64748B")], font_size=8)
            add_styled_cell_text(row[4], [(it.get('status', 'Published'), True, "1E3A8A"), (it.get('awardedDate', '-'), False, "64748B")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        nil = p_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(p_table, [Inches(0.4), Inches(0.65), Inches(3.0), Inches(2.0), Inches(1.15)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 2B. Activities and Initiatives (Exact 6 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("2B. ACTIVITIES AND INITIATIVES (ENTREPRENEURSHIP)", "Startups, incubation, business ideas, and innovation workshops.")
    e_headers = ['S.No', 'Branch', 'Program / Activity Title', 'Date & Organizer', 'Participants', 'Coordinator, Status & Proof']
    e_table = doc.add_table(rows=1, cols=len(e_headers))
    for i, h in enumerate(e_headers): e_table.rows[0].cells[i].text = h

    e_items = gather_items('entrepreneurship')
    if e_items:
        for idx, it in enumerate(e_items):
            row = e_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', ''), ("Type: ", f"{it.get('type', '-')} • {it.get('mode', '-')}", False, "64748B")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('date', '-'), True, "334155"), ("By: ", it.get('organizedBy', '-'), False, "475569")], font_size=8)
            add_styled_cell_text(row[4], [(f"{it.get('participantsCount', '-')} participants", True, "1E3A8A"), (it.get('participants', '-'), False, "64748B")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            
            stat_lines = [("Coord: ", it.get('mentorCoordinator', '-'), True, "1E3A8A"), ("Status: ", it.get('status', 'Completed'), True, "047857")]
            if it.get('link'): stat_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[5], stat_lines, font_size=8)
    else:
        nil = e_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(e_table, [Inches(0.4), Inches(0.65), Inches(2.2), Inches(1.3), Inches(1.15), Inches(1.5)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 3A. Faculty Development Programs (Attended) (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("3A. FACULTY DEVELOPMENT PROGRAMS (FDP) — ATTENDED", "Short-term training and workshops attended by faculty.")
    fa_headers = ['S.No', 'Branch', 'Program Title & Type', 'Organizing Body & Dates', 'Faculty Attended & Proof']
    fa_table = doc.add_table(rows=1, cols=len(fa_headers))
    for i, h in enumerate(fa_headers): fa_table.rows[0].cells[i].text = h

    fa_items = gather_items('fdpAttended', 'fdp')
    if fa_items:
        for idx, it in enumerate(fa_items):
            row = fa_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', ''), ("Type: ", it.get('type', '-'), False, "64748B")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('organizingBody', '-'), False, "334155"), ("Dates: ", f"{it.get('dates', '-')} ({it.get('mode', '-')})", False, "475569")], font_size=8)
            f_lines = [(it.get('facultyAttended', '-'), True, "1E3A8A")]
            if it.get('link'): f_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[4], f_lines, font_size=8)
    else:
        nil = fa_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(fa_table, [Inches(0.4), Inches(0.65), Inches(2.6), Inches(1.8), Inches(1.75)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 3B. Faculty Development Programs (Organized) (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("3B. FACULTY DEVELOPMENT PROGRAMS (FDP) — ORGANIZED", "Faculty development workshops hosted by the institution.")
    fo_headers = ['S.No', 'Branch', 'Program Title & Type', 'Resource Person Details', 'Coordinators & Proof']
    fo_table = doc.add_table(rows=1, cols=len(fo_headers))
    for i, h in enumerate(fo_headers): fo_table.rows[0].cells[i].text = h

    fo_items = gather_items('fdpOrganized')
    if fo_items:
        for idx, it in enumerate(fo_items):
            row = fo_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', ''), ("Type: ", f"{it.get('type', '-')} • {it.get('dates', '-')} ({it.get('mode', '-')})", False, "64748B")], font_size=8.5)
            add_styled_cell_text(row[3], [it.get('resourcePersonDetails', '-')], font_size=8)
            c_lines = [(it.get('facultyCoordinators', '-'), True, "1E3A8A")]
            if it.get('link'): c_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[4], c_lines, font_size=8)
    else:
        nil = fo_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(fo_table, [Inches(0.4), Inches(0.65), Inches(2.3), Inches(2.2), Inches(1.65)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 4. Student Development Programs (SDPs) (Exact 6 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("4. STUDENT DEVELOPMENT PROGRAMS (SDPS)", "Hands-on workshops, guest lectures, internships, and student training.")
    sdp_headers = ['S.No', 'Branch', 'Event Title & Type', 'Date & Resource Person', 'Participants', 'Key Outcomes & Coord.']
    sdp_table = doc.add_table(rows=1, cols=len(sdp_headers))
    for i, h in enumerate(sdp_headers): sdp_table.rows[0].cells[i].text = h

    sdp_items = gather_items('sdp')
    if sdp_items:
        for idx, it in enumerate(sdp_items):
            row = sdp_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', ''), ("Type: ", it.get('type', '-'), False, "64748B")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('date', '-'), True, "334155"), (it.get('resourcePerson', '-'), False, "475569")], font_size=8)
            add_styled_cell_text(row[4], [(f"{it.get('participantsCount', '-')} students", True, "1E3A8A"), (it.get('mode', '-'), False, "64748B")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            
            out_lines = [(it.get('keyOutcomes', '-'), False, "475569")]
            if it.get('coordinator'): out_lines.append(("Coord: ", it.get('coordinator', ''), True, "1E3A8A"))
            if it.get('link'): out_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[5], out_lines, font_size=8)
    else:
        nil = sdp_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(sdp_table, [Inches(0.4), Inches(0.65), Inches(1.9), Inches(1.3), Inches(0.95), Inches(2.0)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 5A. Faculty Achievements (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("5A. FACULTY ACHIEVEMENTS", "Recognitions for research, teaching excellence, and leadership.")
    fac_headers = ['S.No', 'Branch', 'Faculty Member', 'Award / Recognition', 'Conferring Body & Date']
    fac_table = doc.add_table(rows=1, cols=len(fac_headers))
    for i, h in enumerate(fac_headers): fac_table.rows[0].cells[i].text = h

    fac_items = gather_items('facultyAchievements')
    if fac_items:
        for idx, it in enumerate(fac_items):
            row = fac_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('name', '-'), True, "1E3A8A")], font_size=8.5)
            a_lines = [(it.get('award', '-'), True, "0F172A")]
            if it.get('link'): a_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[3], a_lines, font_size=8.5)
            add_styled_cell_text(row[4], [f"{it.get('organization', '-')} ({it.get('date', '-')})"], font_size=8)
    else:
        nil = fac_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(fac_table, [Inches(0.4), Inches(0.65), Inches(1.8), Inches(2.4), Inches(1.95)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 5B. Student Achievements (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("5B. STUDENT ACHIEVEMENTS", "Prizes in competitions, hackathons, symposiums, and sports.")
    stu_headers = ['S.No', 'Branch', 'Student Name & ID', 'Prize / Accomplishment', 'Event & Host Institution']
    stu_table = doc.add_table(rows=1, cols=len(stu_headers))
    for i, h in enumerate(stu_headers): stu_table.rows[0].cells[i].text = h

    stu_items = gather_items('studentAchievements')
    if stu_items:
        for idx, it in enumerate(stu_items):
            row = stu_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('nameRoll', '-'), True, "0F172A")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('award', '-'), True, "0F172A")], font_size=8.5)
            add_styled_cell_text(row[4], [f"{it.get('event', '-')} ({it.get('organization', '-')})"], font_size=8)
    else:
        nil = stu_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(stu_table, [Inches(0.4), Inches(0.65), Inches(2.1), Inches(2.2), Inches(1.85)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 5C. Certifications (Exact 6 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("5C. CERTIFICATIONS", "NPTEL, Coursera, and professional industry certifications.")
    cert_headers = ['S.No', 'Branch', 'Course / Certification Title', 'Platform & Type', 'Duration & Enrolled', 'Certified & Status']
    cert_table = doc.add_table(rows=1, cols=len(cert_headers))
    for i, h in enumerate(cert_headers): cert_table.rows[0].cells[i].text = h

    cert_items = gather_items('certifications')
    if cert_items:
        for idx, it in enumerate(cert_items):
            row = cert_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('title', '')], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('platform', '-'), True, "1E3A8A"), ("Type: ", it.get('type', '-'), False, "64748B")], font_size=8)
            add_styled_cell_text(row[4], [(it.get('duration', '-'), False, "334155"), ("Enrolled: ", it.get('enrolled', '-'), False, "64748B")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[5], [(f"{it.get('certified', '-')} Certified", True, "047857"), (it.get('keyOutcomes', '-'), False, "64748B")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        nil = cert_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(cert_table, [Inches(0.4), Inches(0.65), Inches(2.4), Inches(1.35), Inches(1.15), Inches(1.25)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # Meetings & Minutes (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("MEETINGS & MINUTES", "Official administrative committees, HoD meetings, and governance minutes.")
    m_headers = ['S.No', 'Branch', 'Meeting & Venue', 'Key Decisions & Topics Discussed', 'Policy Directives & Action Plan']
    m_table = doc.add_table(rows=1, cols=len(m_headers))
    for i, h in enumerate(m_headers): m_table.rows[0].cells[i].text = h

    m_items = gather_items('deptMeetings')
    if m_items:
        for idx, it in enumerate(m_items):
            row = m_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            v_lines = [(it.get('title') or "HoD's Meeting", True, "1E3A8A"), (it.get('date', '-'), True, "475569")]
            if it.get('venue'): v_lines.append(("Venue: ", it.get('venue', ''), False, "64748B"))
            add_styled_cell_text(row[2], v_lines, font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

            # Format discussions
            discussions = it.get('discussions', [])
            disc_lines = []
            if it.get('attendees') and len(it.get('attendees', [])) > 0:
                att_names = ", ".join([a.get('name', '') for a in it['attendees'] if isinstance(a, dict)])
                disc_lines.append(("Attendees: ", att_names, False, "1E293B"))
            if discussions and len(discussions) > 0:
                for d in discussions:
                    if isinstance(d, dict):
                        h_txt = d.get('heading', '')
                        d_txt = d.get('details', '').replace('\n', '; ')
                        disc_lines.append((f"{h_txt}: ", d_txt, False, "334155"))
            elif it.get('decisions'):
                disc_lines.append(("", it.get('decisions', ''), False, "334155"))
            add_styled_cell_text(row[3], disc_lines if disc_lines else ["Standard agenda reviewed."], font_size=8)

            pol_lines = [(it.get('policyChanges') or "Standard operations and administrative resolutions confirmed.", False, "475569")]
            if it.get('link'): pol_lines.append(("", it.get('link', ''), False, "2563EB"))
            add_styled_cell_text(row[4], pol_lines, font_size=8)
    else:
        nil = m_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(m_table, [Inches(0.4), Inches(0.65), Inches(1.4), Inches(3.1), Inches(1.65)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 6. Collaborations & MoUs (Exact 6 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("6. COLLABORATIONS & MOUS", "Formal agreements with corporate partners and academic bodies.")
    mou_headers = ['S.No', 'Branch', 'Partner Entity / Industry', 'Date of Signing', 'Faculty SPOC', 'Scope & Focus Area']
    mou_table = doc.add_table(rows=1, cols=len(mou_headers))
    for i, h in enumerate(mou_headers): mou_table.rows[0].cells[i].text = h

    mou_items = gather_items('mous')
    if mou_items:
        for idx, it in enumerate(mou_items):
            row = mou_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('name', '-'), True, "0F172A")], font_size=8.5)
            add_styled_cell_text(row[3], [it.get('datePeriod', '-')], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[4], [(it.get('facultySpoc', '-'), True, "1E3A8A")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[5], [it.get('purpose', '-')], font_size=8)
    else:
        nil = mou_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(mou_table, [Inches(0.4), Inches(0.65), Inches(2.2), Inches(1.15), Inches(1.15), Inches(1.65)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 7. Technical Association Activities (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("7. TECHNICAL ASSOCIATION ACTIVITIES", "Department association workshops, coding leagues, and technical events.")
    tech_headers = ['S.No', 'Branch', 'Association Event & Type', 'Date & Coordinator', 'Participants & Outcomes']
    tech_table = doc.add_table(rows=1, cols=len(tech_headers))
    for i, h in enumerate(tech_headers): tech_table.rows[0].cells[i].text = h

    tech_items = gather_items('techAssociation')
    if tech_items:
        for idx, it in enumerate(tech_items):
            row = tech_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [it.get('event', ''), ("Type: ", it.get('type', '-'), False, "64748B")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('date', '-'), False, "334155"), ("Coord: ", it.get('resourcePersonCoordinator', '-'), True, "1E3A8A")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[4], [(f"{it.get('participants', '-')} participants", True, "0F172A"), (it.get('outcomes', '-'), False, "475569")], font_size=8)
    else:
        nil = tech_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(tech_table, [Inches(0.4), Inches(0.65), Inches(2.3), Inches(1.5), Inches(2.35)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 8. Syllabus Coverage Report (Exact 8 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("8. SYLLABUS COVERAGE REPORT", "Branch-wise academic syllabus delivery status across theory courses and laboratories.")
    syl_headers = ['S.No', 'Branch', 'Course / Subject Title', 'Year/Sem', 'Faculty In-Charge', 'Done (5 Units)', 'Pending / Ongoing', 'Remarks / Status']
    syl_table = doc.add_table(rows=1, cols=len(syl_headers))
    for i, h in enumerate(syl_headers): syl_table.rows[0].cells[i].text = h

    syl_items = gather_items('syllabus')
    if syl_items:
        for idx, it in enumerate(syl_items):
            row = syl_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('subject', '-'), True, "0F172A")], font_size=8)
            add_styled_cell_text(row[3], [normalize_syllabus_year_sem(it.get('yearSem', '-'))], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[4], [(it.get('faculty', '-'), True, "1E3A8A")], font_size=8)
            add_styled_cell_text(row[5], [(normalize_syllabus_completed(it.get('completed'), it.get('pending')), True, "047857")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[6], [(normalize_syllabus_pending(it.get('completed'), it.get('pending')), True, "B45309")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[7], [normalize_syllabus_remarks(it.get('remarks'), it.get('completed'))], font_size=7.5)
    else:
        nil = syl_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(syl_table, [Inches(0.35), Inches(0.55), Inches(1.8), Inches(0.7), Inches(1.2), Inches(0.95), Inches(0.95), Inches(0.7)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 9. Club & Student Engagement Activities (Exact 7 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("9. CLUB & STUDENT ENGAGEMENT ACTIVITIES", "Student club initiatives, campus life programs, and student affairs events.")
    clb_headers = ['S.No', 'Branch', 'Activity / Event Title', 'Type & Duration', 'Participants', 'Coordinator', 'Remarks / Outcomes']
    clb_table = doc.add_table(rows=1, cols=len(clb_headers))
    for i, h in enumerate(clb_headers): clb_table.rows[0].cells[i].text = h

    clb_items = gather_items('studentEngagement')
    if clb_items:
        for idx, it in enumerate(clb_items):
            row = clb_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('title', '-'), True, "0F172A")], font_size=8.5)
            
            type_val = it.get('type', '-')
            dates_val = it.get('dates') or (f"{it.get('startDate', '')} to {it.get('endDate', '')}" if it.get('endDate') else it.get('startDate', '-'))
            add_styled_cell_text(row[3], [(type_val, True, "1E3A8A"), (dates_val, False, "64748B")], font_size=8)
            add_styled_cell_text(row[4], [str(it.get('participantsCount', '-'))], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[5], [(it.get('coordinator', '-'), True, "1E3A8A")], font_size=8)
            add_styled_cell_text(row[6], [it.get('remarks', '-')], font_size=8)
    else:
        nil = clb_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(clb_table, [Inches(0.4), Inches(0.65), Inches(1.9), Inches(1.3), Inches(0.8), Inches(1.1), Inches(1.05)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 10. NSS & Other Extension Activities (Exact 6 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("10. NSS & OTHER EXTENSION ACTIVITIES", "Social outreach camps, community service, and awareness drives.")
    nss_headers = ['S.No', 'Branch', 'Event / Activity Name', 'Date & Venue', 'Participants & Target', 'Outcomes & Coord.']
    nss_table = doc.add_table(rows=1, cols=len(nss_headers))
    for i, h in enumerate(nss_headers): nss_table.rows[0].cells[i].text = h

    nss_items = gather_items('nss')
    if nss_items:
        for idx, it in enumerate(nss_items):
            row = nss_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('event', '-'), True, "0F172A"), ("Type: ", it.get('type', '-'), False, "64748B")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('date', '-'), False, "334155"), ("Venue: ", it.get('venue', '-'), False, "475569")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[4], [(f"{it.get('participantsCount', '-')} participants", True, "1E3A8A"), (it.get('typeOfParticipants', '-'), False, "64748B")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[5], [(it.get('outcomes', '-'), False, "475569"), ("Coord: ", it.get('coordinator', '-'), True, "1E3A8A")], font_size=8)
    else:
        nil = nss_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(nss_table, [Inches(0.4), Inches(0.65), Inches(1.8), Inches(1.3), Inches(1.2), Inches(1.85)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 11. Additional/Other Relevant Initiatives (Exact 5 columns as PDF)
    # -------------------------------------------------------------
    add_major_heading("11. ADDITIONAL/OTHER RELEVANT INITIATIVES", "Special institutional initiatives, lab modernizations, and quality improvements.")
    add_headers = ['S.No', 'Branch', 'Initiative / Activity', 'Date & Coordinator', 'Description & Key Outcomes']
    add_table = doc.add_table(rows=1, cols=len(add_headers))
    for i, h in enumerate(add_headers): add_table.rows[0].cells[i].text = h

    add_items = gather_items('additionalInitiatives')
    if add_items:
        for idx, it in enumerate(add_items):
            row = add_table.add_row().cells
            row[0].text = str(idx + 1); row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row[1].text = it['deptCode']; row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; row[1].paragraphs[0].runs[0].bold = True

            add_styled_cell_text(row[2], [(it.get('initiative', '-'), True, "0F172A")], font_size=8.5)
            add_styled_cell_text(row[3], [(it.get('date', '-'), False, "334155"), ("Coord: ", it.get('coordinator', '-'), True, "1E3A8A")], font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
            add_styled_cell_text(row[4], [(it.get('description', '-'), False, "475569"), ("Outcome: ", it.get('outcomes', '-'), True, "047857")], font_size=8)
    else:
        nil = add_table.add_row().cells
        for c in nil: c.text = "NIL"; c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_table_styling(add_table, [Inches(0.4), Inches(0.65), Inches(2.4), Inches(1.35), Inches(2.4)])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)

    return {
        'success': True,
        'output_path': output_path,
        'period': period
    }

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_consolidated_master_docx.py <input_json_file_or_string> <output_docx_path>")
        sys.exit(1)

    input_source = sys.argv[1]
    output_docx = sys.argv[2]

    if os.path.exists(input_source):
        with open(input_source, 'r', encoding='utf-8') as f:
            data = json.load(f)
    else:
        data = json.loads(input_source)

    period_val = data.get('period', 'September 2026')
    depts_data = data.get('departmentsData', {})

    res = generate_consolidated_master_docx(depts_data, period_val, output_docx)
    print(json.dumps(res))
